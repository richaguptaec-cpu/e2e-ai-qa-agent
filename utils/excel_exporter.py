from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def create_excel_report(
    validated_suite
):

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Test Cases"

    headers = [
        "Test ID",
        "Title",
        "Type",
        "Priority",
        "Source Agent",
        "Quality Score",
        "Quality Status",
        "Validation Status",
        "Precondition",
        "Steps",
        "Expected Result"
    ]

    worksheet.append(headers)

    # --------------------------------------------------------
    # Header formatting
    # --------------------------------------------------------

    for cell in worksheet[1]:

        cell.font = Font(
            bold=True
        )

    # --------------------------------------------------------
    # Test Cases
    # --------------------------------------------------------

    for test in validated_suite.test_cases:

        steps = "\n".join(
            f"{index}. {step}"
            for index, step in enumerate(
                test.steps,
                start=1
            )
        )

        worksheet.append([
            test.test_id,
            test.title,
            test.type,
            test.priority,
            test.source_agent,
            test.quality_score,
            test.quality_status,
            test.validation_status,
            test.precondition,
            steps,
            test.expected_result
        ])

    # --------------------------------------------------------
    # Column widths
    # --------------------------------------------------------

    widths = {
        1: 12,
        2: 45,
        3: 25,
        4: 15,
        5: 25,
        6: 15,
        7: 20,
        8: 20,
        9: 45,
        10: 70,
        11: 60
    }

    for column, width in widths.items():

        worksheet.column_dimensions[
            get_column_letter(column)
        ].width = width

    # --------------------------------------------------------
    # Text wrapping
    # --------------------------------------------------------

    for row in worksheet.iter_rows():

        for cell in row:

            cell.alignment = cell.alignment.copy(
                wrap_text=True,
                vertical="top"
            )

    # --------------------------------------------------------
    # Freeze header
    # --------------------------------------------------------

    worksheet.freeze_panes = "A2"

    # --------------------------------------------------------
    # Save in memory
    # --------------------------------------------------------

    output = BytesIO()

    workbook.save(
        output
    )

    output.seek(0)

    return output
